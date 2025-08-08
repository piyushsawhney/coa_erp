import os
import time

from openpyxl import Workbook, load_workbook
from selenium.common import NoSuchElementException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.wait import WebDriverWait

from bni.selenium_setup import driver

data = []
file_path = "members.xlsx"
headers = ["Name", "Profile Link", "Company", "Profession", "Mobile", "Phone", "Email"]

if not os.path.exists(file_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)  # write headers
    wb.save(file_path)


def append_row(data_dict):
    wb = load_workbook(file_path)
    ws = wb.active
    row = [data_dict.get(h, "") for h in headers]
    ws.append(row)
    wb.save(file_path)


def get_email_link_from_profile():
    email_link = None
    ul_blocks = driver.find_elements(By.CSS_SELECTOR, "ul.memberContactInfo")
    for ul in ul_blocks:
        a_tags = ul.find_elements(By.TAG_NAME, "a")
        for a in a_tags:
            href = a.get_attribute("href")
            if href and "sendmessage" in href:
                return href
    return email_link


def get_mobile_from_profile():
    element = driver.find_element(By.CSS_SELECTOR, "a.moredots")
    driver.execute_script("arguments[0].scrollIntoView({ behavior: 'smooth', block: 'center' });", element)
    element.click()
    time.sleep(1)
    phone_number = None
    direct_number = None

    try:
        phone_element = driver.find_element(By.XPATH, "//li[strong[text()='Phone']]/a")
        phone_number = phone_element.text.strip()
    except NoSuchElementException:
        pass

    try:
        direct_element = driver.find_element(By.XPATH, "//li[strong[text()='Direct']]/a")
        direct_number = direct_element.text.strip()
    except NoSuchElementException:
        pass
    return phone_number, direct_number


def get_email_mobile_from_profile():
    WebDriverWait(
        driver,
        timeout=15,
        poll_frequency=1,
        ignored_exceptions=[NoSuchElementException]
    ).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "a.moredots")))
    email_link = get_email_link_from_profile()
    phone1, phone2 = get_mobile_from_profile()
    return phone1, phone2, email_link


def navigate_to_members_tab_and_click():
    member_tab_id = "members_tab"
    member_tab_element = WebDriverWait(
        driver,
        timeout=15,
        poll_frequency=1,
        ignored_exceptions=[NoSuchElementException]
    ).until(EC.visibility_of_element_located((By.ID, member_tab_id)))
    driver.execute_script("arguments[0].scrollIntoView({ behavior: 'smooth', block: 'center' });", member_tab_element)
    driver.execute_script("window.scrollBy(0, 800);")
    member_tab_button = WebDriverWait(
        driver,
        timeout=15,
        poll_frequency=1,
        ignored_exceptions=[NoSuchElementException]
    ).until(EC.element_to_be_clickable((By.ID, member_tab_id)))
    member_tab_button.click()


def get_member_list():
    return WebDriverWait(
        driver,
        timeout=15,
        poll_frequency=1,
        ignored_exceptions=[NoSuchElementException]
    ).until(EC.visibility_of_all_elements_located((By.CSS_SELECTOR, "#chapterListTable tbody tr")))


def navigate_members_in_a_table_and_process_data():
    total_rows = get_member_list()
    row_count = len(total_rows)
    for i in range(row_count):
        row = get_member_list()[i]
        tds = row.find_elements(By.TAG_NAME, "td")
        if len(tds) >= 3:
            profile_link_tag = WebDriverWait(tds[0], 10).until(
                EC.element_to_be_clickable((By.TAG_NAME, "a"))
            )
            profile_url = profile_link_tag.get_attribute("href")
            member_name = profile_link_tag.text.strip()
            company_name = tds[1].text.strip()
            profession = tds[2].text.strip()
            driver.execute_script("arguments[0].scrollIntoView({behavior: 'smooth', block: 'center'});",
                                  profile_link_tag)
            time.sleep(1)
            driver.execute_script("window.scrollBy(0, 600);")
            profile_link_tag.click()
            phone1, phone2, email = get_email_mobile_from_profile()
            data_dict = {
                "Name": member_name,
                "Profile Link": profile_url,
                "Company": company_name,
                "Profession": profession,
                "Mobile": phone1,
                "Phone": phone2,
                "Email": email
            }
            append_row(data_dict)
            driver.back()
            navigate_to_members_tab_and_click()


def process_member_navigation():
    navigate_to_members_tab_and_click()
    navigate_members_in_a_table_and_process_data()
    return 0
